"""Generate comprehensive Word (.docx) and PDF documentation for SkillProof."""
import os
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# ══════════════════════════════════════════════════════════════════════════════
# 1. DOCX GENERATION
# ══════════════════════════════════════════════════════════════════════════════
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_docx(filepath):
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("SkillProof Platform")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(99, 102, 241) # Indigo

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(20)
    run_sub = sub_p.add_run("Project-Based Career Evidence & Skill Verification System\nComprehensive Technical Architecture & Execution Manual")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(107, 114, 128)

    # Metadata table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Name:", "SkillProof (Beyond Certificates — Show What You Can Actually Do)"),
        ("Architecture:", "FastAPI (Python 3.11) + SQLAlchemy ORM + SQLite + Vanilla Web App"),
        ("Date / Release:", "October 2026 | Version 1.0.0 Production Release"),
        ("Documentation Scope:", "System Architecture, Scoring Engine, API Specs, Personas & Execution"),
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = k
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c1.text = v
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Section 1
    h1 = doc.add_heading("1. Executive Summary & Problem Statement", level=1)
    h1.paragraph_format.space_before = Pt(16)
    
    p = doc.add_paragraph()
    p.add_run("In modern software engineering recruiting, traditional resumes and certificates of completion have suffered severe credential inflation. Anyone can claim proficiency in Python, React, or Distributed Systems by listing keywords on a resume or finishing a passive video tutorial course. Recruiters and hiring managers spend hundreds of hours filtering through inflated resumes, while talented students struggle to differentiate themselves.")

    p2 = doc.add_paragraph()
    p2.add_run("SkillProof introduces a paradigm shift: ")
    r_bold = p2.add_run("Project-Based Career Evidence.")
    r_bold.bold = True
    p2.add_run(" Instead of relying on self-reported assertions, SkillProof ingests a student's public GitHub repositories, analyzes the raw commit timeline, evaluates repository complexity, assesses documentation quality, and computes a transparent, deterministic Evidence Score (0–100). Furthermore, it bridges automated analysis with human accountability through mentor and faculty verification.")

    # Section 2
    h2 = doc.add_heading("2. High-Level System Architecture", level=1)
    
    p3 = doc.add_paragraph()
    p3.add_run("SkillProof is architected with a decoupled, high-performance client-server model:")

    arch_bullets = [
        ("Presentation Layer (Frontend):", "Single-Page Application (SPA) built using semantic HTML5, custom responsive CSS3 design tokens with dark-mode glassmorphism, and vanilla ES6 JavaScript. Visualizes commit activity using Chart.js and vector iconography via Lucide."),
        ("Application & API Layer (Backend):", "FastAPI web framework providing asynchronous endpoints, Pydantic v2 request/response validation schemas, and automated OpenAPI (Swagger) documentation."),
        ("Data Persistence Layer:", "SQLAlchemy ORM with SQLite database engine. Structured with canonical path resolution to ensure multi-process consistency across seeds, tests, and web servers."),
        ("Deterministic Scoring Engine:", "A multi-factor mathematical scoring pipeline with diminishing-returns curves, consistency ratios, and anti-gaming penalties."),
        ("AI Summary & Skills Extraction:", "A dual-mode intelligence service that extracts skills and plain-language summaries using NLP heuristic rule sets with seamless OpenAI API integration when configured."),
        ("GitHub Integration Pipeline:", "Asynchronous HTTP client leveraging the GitHub REST API v3 to inspect commits, language byte distributions, file counts, and README documents."),
    ]
    for b_title, b_desc in arch_bullets:
        bp = doc.add_paragraph(style='List Bullet')
        r = bp.add_run(b_title + " ")
        r.bold = True
        bp.add_run(b_desc)

    # Section 3: Scoring Engine
    h3 = doc.add_heading("3. The Evidence Scoring Engine & Anti-Gaming", level=1)
    p_sc = doc.add_paragraph()
    p_sc.add_run("The core value of SkillProof is its objective, transparent scoring engine. Each repository is evaluated across five pillars totaling 100 points, followed by anti-gaming audits:")

    # Scoring table
    st_table = doc.add_table(rows=6, cols=3)
    st_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    st_headers = ["Pillar", "Weight", "Scoring Criteria & Formula"]
    for j, h in enumerate(st_headers):
        cell = st_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "E0E7FF")

    st_rows = [
        ("Commit Consistency", "30 pts", "Rewards sustained activity across multiple months. Ratio of active months to total project lifespan. Base score (up to 15 pts) + consistency bonus (up to 15 pts)."),
        ("Project Complexity", "25 pts", "Multi-language balance (up to 10 pts for 3+ languages), file count depth (up to 8 pts for 30+ files), and code size (up to 7 pts for >500KB)."),
        ("README Quality", "20 pts", "Length (up to 10 pts for >1500 chars) + quality markers (install instructions, usage guides, markdown headers, code blocks, contributing guides, licensing)."),
        ("Recent Activity", "15 pts", "Recency of last commit. Full points if active within 30 days; scaled down for older projects (>1 year = 2 pts)."),
        ("Community Signal", "10 pts", "Repository stars (up to 6 pts) and originality verification (original repo vs active fork)."),
    ]
    for i, row in enumerate(st_rows, start=1):
        for j, text in enumerate(row):
            c = st_table.rows[i].cells[j]
            c.text = text
            c.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    p_anti = doc.add_paragraph()
    r_aw = p_anti.add_run("Anti-Gaming Audits & Penalties:")
    r_aw.bold = True
    anti_list = [
        ("Single-Day Commit Dump (-15 pts):", "Flags repositories where commit timestamps cluster unnaturally in a single 24-hour window, discouraging last-minute fake commits."),
        ("Bare Fork Detection (-30 pts / capped at 10):", "Detects forked repositories that contain zero original commit contributions from the student."),
        ("Minimal Contribution Warning (-10 pts):", "Flags repositories with negligible code size or single-file trivial edits."),
    ]
    for a_title, a_desc in anti_list:
        ap = doc.add_paragraph(style='List Bullet')
        r = ap.add_run(a_title + " ")
        r.bold = True
        ap.add_run(a_desc)

    # Section 4: Data Models
    doc.add_heading("4. Database Schema & Data Models", level=1)
    models_info = [
        ("Student (`students`)", "id, name, email, github_username, bio, avatar_url, evidence_score, score_breakdown (JSON), created_at, updated_at."),
        ("Project (`projects`)", "id, student_id, github_repo_url, name, description, languages (JSON), topics (JSON), commit_count, commit_timeline (JSON), readme_content, is_fork, star_count, is_public, ai_summary, ai_skills (JSON), evidence_score, score_breakdown (JSON), anti_gaming_flags (JSON)."),
        ("Skill (`skills`)", "id, student_id, name, level (beginner/intermediate/advanced), source ('ai-suggested' or 'verified'), verified_by, verified_at, evidence_project_ids (JSON), monthly_activity (JSON)."),
        ("MentorFeedback (`mentor_feedback`)", "id, project_id, mentor_name, mentor_email, rating (1-5), strengths, improvements, comments, skills_verified (JSON), created_at."),
    ]
    for m_name, m_fields in models_info:
        mp = doc.add_paragraph()
        r = mp.add_run(m_name + ": ")
        r.bold = True
        mp.add_run(m_fields)

    # Section 5: API Endpoints
    doc.add_heading("5. RESTful API Specification", level=1)
    api_table = doc.add_table(rows=9, cols=3)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Method & Route", "Description", "Key Parameters / Payloads"]
    for j, h in enumerate(api_headers):
        cell = api_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "F3F4F6")

    api_rows = [
        ("GET /api/health", "System health check", "None"),
        ("GET /api/students/", "List all students with brief summary", "None"),
        ("POST /api/students/", "Create new student profile", "name, email, github_username, bio"),
        ("GET /api/students/{id}", "Full candidate portfolio & proof audit", "Path: student_id"),
        ("PUT /api/students/{s_id}/projects/{p_id}/visibility", "Toggle project Public/Private privacy", "is_public: boolean"),
        ("GET /api/recruiter/search", "Filter candidates by skill & minimum score", "skill, min_score (0-100), sort_by"),
        ("GET /api/recruiter/skills", "Unique skills catalog across candidate pool", "None"),
        ("POST /api/feedback/", "Submit mentor review & verify skills", "project_id, mentor_name, rating, skills_verified"),
    ]
    for i, row in enumerate(api_rows, start=1):
        for j, text in enumerate(row):
            c = api_table.rows[i].cells[j]
            c.text = text
            c.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 6: Personas & User Journeys
    doc.add_heading("6. Key User Personas & Workflows", level=1)
    p_pj = doc.add_paragraph()
    p_pj.add_run("1. The Student / Job Seeker: Connects their public GitHub username to generate an evidence-based profile. Can selectively hide sensitive projects (privacy toggle), review detailed score breakdowns, and share their verifiable portfolio URL with recruiters.\n")
    p_pj.add_run("2. The Technical Recruiter / Hiring Manager: Searches candidate pools filtered by required stack (e.g. Flutter, Docker, Python) and dials in minimum evidence thresholds to instantly discover proven builders rather than keyword optimizers.\n")
    p_pj.add_run("3. The Mentor / University Faculty: Inspects student codebases, writes constructive qualitative feedback, assigns 1-5 star ratings, and checks boxes to convert AI-suggested skills into permanent human-verified credentials.")

    # Section 7: Execution Guide Summary
    doc.add_heading("7. Execution & Running the Project", level=1)
    p_run = doc.add_paragraph()
    p_run.add_run("The complete instructions are preserved in ")
    p_run.add_run("HOW_TO_RUN.txt").bold = True
    p_run.add_run(". Quick reference commands:\n\n")
    p_run.add_run("1. Setup Environment:  pip install -r backend/requirements.txt\n")
    p_run.add_run("2. Seed Database:      python -m app.seed (loads 5 demo student profiles)\n")
    p_run.add_run("3. Run Test Suite:     pytest backend/tests -v (13/13 tests pass)\n")
    p_run.add_run("4. Launch Web Server:  uvicorn app.main:app --host 127.0.0.1 --port 8000\n")
    p_run.add_run("5. Open Application:   http://127.0.0.1:8000/\n")

    doc.save(filepath)
    print(f"[OK] Word document generated at: {filepath}")


# ══════════════════════════════════════════════════════════════════════════════
# 2. PDF GENERATION (ReportLab)
# ══════════════════════════════════════════════════════════════════════════════
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#6b7280"))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, 750, "SkillProof — Project-Based Career Evidence Platform")
            self.drawRightString(612 - 54, 750, "Technical Documentation & Execution Guide")
            self.setStrokeColor(colors.HexColor("#e5e7eb"))
            self.setLineWidth(0.5)
            self.line(54, 742, 612 - 54, 742)

        # Footer
        self.setStrokeColor(colors.HexColor("#e5e7eb"))
        self.setLineWidth(0.5)
        self.line(54, 45, 612 - 54, 45)
        self.drawString(54, 32, "Confidential & Proprietary — SkillProof System Architecture")
        self.drawRightString(612 - 54, 32, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


def generate_pdf(filepath):
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor("#4338ca"),
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#4b5563"),
        spaceAfter=15,
    )
    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#1e1b4b"),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True,
    )
    h2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#3730a3"),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#1f2937"),
        spaceAfter=8,
    )
    bullet_style = ParagraphStyle(
        'Bullet',
        parent=body_style,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=4,
    )
    code_style = ParagraphStyle(
        'CodeBlock',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1e293b"),
        backColor=colors.HexColor("#f1f5f9"),
        borderPadding=6,
        spaceBefore=4,
        spaceAfter=8,
    )

    story = []

    # Title Banner
    story.append(Paragraph("SkillProof Platform", title_style))
    story.append(Paragraph("Project-Based Career Evidence & Skill Verification System<br/><b>Technical Specification, Engine Formulas, and Execution Guide</b>", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#6366f1"), spaceAfter=12))

    # Overview Table
    meta_table_data = [
        [Paragraph("<b>Release:</b> v1.0.0 Production", body_style), Paragraph("<b>Tech Stack:</b> FastAPI / SQLite / ES6", body_style)],
        [Paragraph("<b>Author:</b> SkillProof Architecture Team", body_style), Paragraph("<b>Verification:</b> 13/13 Automated Tests Pass", body_style)],
        [Paragraph("<b>App URL:</b> http://127.0.0.1:8000/", body_style), Paragraph("<b>API Docs:</b> http://127.0.0.1:8000/docs", body_style)],
    ]
    meta_table = Table(meta_table_data, colWidths=[240, 264])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # Section 1
    story.append(Paragraph("1. Executive Summary & Problem Statement", h1_style))
    story.append(Paragraph(
        "Modern technical hiring is broken by credential inflation. Traditional resumes, keyword lists, and course certificates of completion fail to reflect authentic software engineering capability. Recruiters are inundated with applicants claiming identical skill sets, while talented students struggle to substantiate their competence.",
        body_style
    ))
    story.append(Paragraph(
        "<b>SkillProof</b> solves this through <b>Project-Based Career Evidence</b>. By ingesting public GitHub repositories and analyzing actual commit cadence, codebase complexity, language distribution, and documentation quality, SkillProof produces a transparent <b>0–100 Evidence Score</b>. Automated insights are coupled with faculty/mentor peer review to deliver tamper-resistant skill verification.",
        body_style
    ))

    # Section 2
    story.append(Paragraph("2. System Architecture & Component Design", h1_style))
    arch_items = [
        "<b>Presentation Layer:</b> Single-Page Application using HTML5, modern CSS3 dark-mode glassmorphism, and responsive vanilla ES6 JavaScript. Visualizes commit activity using Chart.js line charts and vector icons via Lucide.",
        "<b>Application / API Server:</b> FastAPI asynchronous microframework handling RESTful routing, input validation with Pydantic v2 schemas, and interactive OpenAPI documentation.",
        "<b>Data Persistence Layer:</b> SQLite database backed by SQLAlchemy ORM with canonical absolute path resolution for reliable multi-process operations.",
        "<b>Deterministic Scoring Engine:</b> Multi-variable mathematical pipeline evaluating commit consistency, repo complexity, README quality, recent activity, and anti-gaming anomalies.",
        "<b>AI Summary & Skill Extraction:</b> Hybrid intelligent parser extracting skills from language profiles, topics, and README keywords with OpenAI API support and rule-based fallback.",
    ]
    for it in arch_items:
        story.append(Paragraph(f"• {it}", bullet_style))

    story.append(Spacer(1, 8))

    # Section 3: Scoring Table
    story.append(Paragraph("3. Deterministic Evidence Scoring Engine", h1_style))
    story.append(Paragraph(
        "The Evidence Score evaluates repos across 5 key pillars (max 100 pts) followed by anti-gaming penalty subtractions:",
        body_style
    ))

    score_data = [
        [Paragraph("<b>Pillar</b>", body_style), Paragraph("<b>Max Pts</b>", body_style), Paragraph("<b>Mathematical Logic & Criteria</b>", body_style)],
        [Paragraph("Commit Consistency", body_style), Paragraph("30 pts", body_style), Paragraph("Evaluates sustained development across months. Diminishing returns on commit count (up to 15 pts) + consistency ratio across active months (up to 15 pts).", body_style)],
        [Paragraph("Project Complexity", body_style), Paragraph("25 pts", body_style), Paragraph("Language distribution balance (up to 10 pts for 3+ languages), file count depth (up to 8 pts for 30+ files), and code size (up to 7 pts for >500KB).", body_style)],
        [Paragraph("README Quality", body_style), Paragraph("20 pts", body_style), Paragraph("Length score (up to 10 pts for >1500 chars) + quality markers (install, usage, code blocks, headers, license, contributing).", body_style)],
        [Paragraph("Recent Activity", body_style), Paragraph("15 pts", body_style), Paragraph("Recency of last commit: 15 pts if <30 days; 10 pts if <90 days; 6 pts if <180 days; 2 pts if >1 year.", body_style)],
        [Paragraph("Community Signal", body_style), Paragraph("10 pts", body_style), Paragraph("Repository stars (up to 6 pts) and originality verification (original repo vs original commits in fork).", body_style)],
    ]
    score_table = Table(score_data, colWidths=[120, 55, 329])
    score_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#e0e7ff")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#94a3b8")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(score_table)
    story.append(Spacer(1, 8))

    story.append(Paragraph("<b>Anti-Gaming Safeguards:</b>", h2_style))
    anti_items = [
        "<b>Single-Day Burst Penalty (-15 pts):</b> Penalizes repos where all commits were dumped in a single 24h period.",
        "<b>Bare Fork Penalty (-30 pts, capped at 10):</b> Penalizes forks that contain zero original commit contributions.",
        "<b>Student Score Aggregation:</b> Computes a weighted average favoring the top 3 best projects, plus multi-project quality bonuses (up to +10 pts) for students with multiple verified projects.",
    ]
    for it in anti_items:
        story.append(Paragraph(f"• {it}", bullet_style))

    story.append(PageBreak())

    # Section 4: Data Models & APIs
    story.append(Paragraph("4. Database Models & Schema Design", h1_style))
    models_text = [
        "<b>Student:</b> Stores user profile, email, GitHub handle, evidence score, and JSON score breakdown.",
        "<b>Project:</b> Stores repo metadata, languages byte dict, commit timeline array, README content, AI summary, and anti-gaming audit flags.",
        "<b>Skill:</b> Stores skill name, proficiency level, evidence source (AI-suggested vs verified), and verifying mentor name.",
        "<b>MentorFeedback:</b> Stores project feedback, 1-5 star ratings, qualitative strengths/improvements, and verified skills.",
    ]
    for it in models_text:
        story.append(Paragraph(f"• {it}", bullet_style))

    story.append(Spacer(1, 8))
    story.append(Paragraph("5. API Route Specifications", h1_style))

    api_data = [
        [Paragraph("<b>Endpoint</b>", body_style), Paragraph("<b>Method</b>", body_style), Paragraph("<b>Function & Purpose</b>", body_style)],
        [Paragraph("/api/health", body_style), Paragraph("GET", body_style), Paragraph("Service health check status confirmation", body_style)],
        [Paragraph("/api/students/", body_style), Paragraph("GET", body_style), Paragraph("List all students with brief summary & top skills", body_style)],
        [Paragraph("/api/students/{id}", body_style), Paragraph("GET", body_style), Paragraph("Full student portfolio, commit timelines, and audit", body_style)],
        [Paragraph("/api/students/{s_id}/projects/{p_id}/visibility", body_style), Paragraph("PUT", body_style), Paragraph("Toggle project visibility (Public vs Private privacy control)", body_style)],
        [Paragraph("/api/recruiter/search", body_style), Paragraph("GET", body_style), Paragraph("Filter candidates by skill name and min score slider", body_style)],
        [Paragraph("/api/recruiter/skills", body_style), Paragraph("GET", body_style), Paragraph("List all unique skills with total and verified counts", body_style)],
        [Paragraph("/api/feedback/", body_style), Paragraph("POST", body_style), Paragraph("Submit mentor review & verify candidate skills", body_style)],
    ]
    api_table = Table(api_data, colWidths=[180, 50, 274])
    api_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#f1f5f9")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#94a3b8")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(api_table)
    story.append(Spacer(1, 10))

    # Section 6: Personas
    story.append(Paragraph("6. User Personas & Workflows", h1_style))
    story.append(Paragraph(
        "<b>Student:</b> Imports GitHub repositories, visualizes score breakdown, toggles project privacy, and shares verifiable link.<br/>"
        "<b>Recruiter:</b> Searches candidate talent pool by verified skill badges and minimum evidence score filter.<br/>"
        "<b>Mentor:</b> Conducts peer review, rates projects 1-5 stars, writes qualitative evaluations, and certifies skills.",
        body_style
    ))

    # Section 7: Running the Project
    story.append(Paragraph("7. Quick Execution Guide", h1_style))
    story.append(Paragraph("Refer to <b>HOW_TO_RUN.txt</b> for comprehensive instructions. Key execution steps:", body_style))
    
    code_snippet = (
        "# 1. Install dependencies\n"
        "pip install -r backend/requirements.txt\n\n"
        "# 2. Seed database with demo candidates\n"
        "python -m app.seed\n\n"
        "# 3. Run automated tests (13/13 pass)\n"
        "pytest backend/tests -v\n\n"
        "# 4. Launch web application server\n"
        "uvicorn app.main:app --host 127.0.0.1 --port 8000\n\n"
        "# 5. Access in browser: http://127.0.0.1:8000/"
    )
    story.append(Paragraph(code_snippet.replace("\n", "<br/>"), code_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[OK] PDF document generated at: {filepath}")


if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(out_dir, "SkillProof_Project_Documentation.docx")
    pdf_path = os.path.join(out_dir, "SkillProof_Project_Documentation.pdf")
    
    generate_docx(docx_path)
    generate_pdf(pdf_path)
    print("All documents generated successfully.")
