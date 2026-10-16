"""Generate professional 16:9 widescreen PowerPoint deck for SkillProof project presentation."""
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="SkillProof_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    BG_LIGHT = RGBColor(248, 250, 252)   # Slate 50
    TEXT_DARK = RGBColor(15, 23, 42)     # Slate 900
    TEXT_LIGHT = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    PRIMARY = RGBColor(67, 56, 202)      # Indigo 700
    PRIMARY_LIGHT = RGBColor(99, 102, 241) # Indigo 500
    ACCENT_GREEN = RGBColor(16, 185, 129) # Emerald 500
    ACCENT_AMBER = RGBColor(245, 158, 11) # Amber 500
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240)
    CARD_BG_TINT = RGBColor(241, 245, 249)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, section_tag, title_text, dark=False):
        # Section tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = section_tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = PRIMARY_LIGHT if dark else PRIMARY
        
        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT if dark else TEXT_DARK

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    # Pill badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(3.8), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 41, 59)
    badge.line.color.rgb = PRIMARY_LIGHT
    tb = badge.text_frame
    tb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tb.paragraphs[0]
    p.text = "⚡ NEXT-GEN CAREER EVIDENCE PLATFORM"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_LIGHT
    p.alignment = PP_ALIGN.CENTER

    # Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(1.6))
    tf = t_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "SkillProof"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "Beyond Resumes & Certificates — Show What You Can Actually Do"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_before = Pt(8)

    # 3 Stat Cards on Title
    card_data = [
        ("Deterministic Evidence", "0–100 Evidence Score based on real git timeline, code volume & docs"),
        ("Dual-Mode Verification", "Automated GitHub commit parsing + Human faculty/mentor sign-off"),
        ("Anti-Gaming Engine", "Penalizes burst commit dumps, bare forks & keyword stuffing"),
    ]
    for i, (head, sub) in enumerate(card_data):
        c_left = Inches(1.0 + i * 3.85)
        add_card(s1, c_left, Inches(4.2), Inches(3.6), Inches(1.8), bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
        c_tb = s1.shapes.add_textbox(c_left + Inches(0.2), Inches(4.4), Inches(3.2), Inches(1.4))
        c_tf = c_tb.text_frame
        c_tf.word_wrap = True
        cp1 = c_tf.paragraphs[0]
        cp1.text = head
        cp1.font.size = Pt(14)
        cp1.font.bold = True
        cp1.font.color.rgb = RGBColor(224, 231, 255)
        cp2 = c_tf.add_paragraph()
        cp2.text = sub
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = RGBColor(148, 163, 184)
        cp2.space_before = Pt(6)

    # Footer note
    f_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.5))
    f_tf = f_box.text_frame
    fp = f_tf.paragraphs[0]
    fp.text = "Presentation Duration: 7 Minutes  |  Full-Stack Production Architecture: FastAPI + SQLite + SPA + WebSocket"
    fp.font.size = Pt(11)
    fp.font.color.rgb = RGBColor(100, 116, 139)

    # =========================================================================
    # SLIDE 2: 1. The Problem
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Problem Definition", "The Crisis of Credential Inflation & Unverifiable Claims")

    p_cols = [
        ("Target User", [
            ("Computer Science & Engineering Students", "Struggling to stand out with zero prior corporate work history."),
            ("Early-Career Developers", "Building impressive personal projects that get overlooked in automated ATS filters."),
            ("Technical Recruiters & Hiring Managers", "Drowning in 500+ identical resumes per open software role."),
            ("Academic Mentors & Bootcamps", "Seeking reliable ways to endorse and vouch for their students' true technical skills.")
        ]),
        ("Current Process & Failures", [
            ("Keyword-Stuffed Resumes", "Candidates list Python, Docker, React without having written more than 50 lines."),
            ("Passive Course Certificates", "Certificates of completion only prove video playback, not coding capability."),
            ("Manual GitHub Auditing", "Recruiters lack the technical time or tools to dive into commit logs of 200 candidates."),
            ("Costly Technical Screens", "Companies waste 40+ engineering hours interviewing candidates who cannot build.")
        ]),
        ("Main Problem & Why It Matters", [
            ("The Signal-to-Noise Deficit", "Hiring teams cannot distinguish genuine software builders from tutorial copycats."),
            ("False Positives / Bad Hires", "Average cost of a bad technical hire exceeds $50,000 in onboarding and lost velocity."),
            ("Demoralized Talent", "High-capability students from non-target universities get ignored because of resume pedigree."),
            ("Why It Matters Now", "With AI code generation, anyone can generate boilerplate; proof of sustained effort is vital.")
        ])
    ]

    for i, (title, items) in enumerate(p_cols):
        left = Inches(0.8 + i * 3.95)
        add_card(s2, left, Inches(1.6), Inches(3.8), Inches(5.3), bg_color=CARD_BG, border_color=CARD_BORDER)
        
        # Header inside card
        tb = s2.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(3.4), Inches(0.6))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        # Items
        tb_body = s2.shapes.add_textbox(left + Inches(0.2), Inches(2.4), Inches(3.4), Inches(4.3))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        for idx, (subhead, text) in enumerate(items):
            bp = tf_body.add_paragraph() if idx > 0 else tf_body.paragraphs[0]
            bp.text = f"• {subhead}: "
            bp.font.size = Pt(10.5)
            bp.font.bold = True
            bp.font.color.rgb = TEXT_DARK
            bp.space_before = Pt(8) if idx > 0 else Pt(0)
            
            run = bp.add_run()
            run.text = text
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: 2. The Solution
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Solution Overview", "SkillProof: Transforming Codebases into Verifiable Career Evidence")

    # Top summary card
    add_card(s3, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.0), bg_color=RGBColor(238, 242, 255), border_color=PRIMARY_LIGHT)
    tb_top = s3.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(0.8))
    tf_top = tb_top.text_frame
    tf_top.word_wrap = True
    pt = tf_top.paragraphs[0]
    pt.text = "SOLUTION OVERVIEW: An autonomous, transparent platform that turns raw Git commit histories, repository architectures, and documentation quality into an objective 0–100 Evidence Score backed by mentor validation."
    pt.font.size = Pt(12)
    pt.font.bold = True
    pt.font.color.rgb = PRIMARY

    # 4 Key Feature Cards
    sol_features = [
        ("1. Git Commit Ingestion", "Deep inspection of public repositories via GitHub REST API. Tracks multi-month commit cadence, language byte ratios, file volume, and README markdown depth."),
        ("2. Deterministic Scoring", "Mathematical 5-pillar formula (100 pts) evaluating Consistency, Complexity, Documentation, Recency, and Community signals with full audit explainability."),
        ("3. Anti-Gaming Audits", "Active heuristic algorithms that flag single-day commit dumps (-15 pts), bare forks with no original commits (-30 pts), and trivial repository shells."),
        ("4. Human-in-the-Loop", "Mentors and faculty review projects, score technical depth (1-5★), and convert AI-suggested candidate skills into permanently verified credentials.")
    ]
    for i, (f_title, f_desc) in enumerate(sol_features):
        left = Inches(0.8 + i * 2.95)
        add_card(s3, left, Inches(2.7), Inches(2.8), Inches(2.2), bg_color=CARD_BG, border_color=CARD_BORDER)
        tb_f = s3.shapes.add_textbox(left + Inches(0.15), Inches(2.85), Inches(2.5), Inches(1.9))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True
        p = tf_f.paragraphs[0]
        p.text = f_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p2 = tf_f.add_paragraph()
        p2.text = f_desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(6)

    # Before vs After Comparison Card
    add_card(s3, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.9), bg_color=CARD_BG_TINT, border_color=CARD_BORDER)
    tb_comp = s3.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.3), Inches(1.7))
    tf_comp = tb_comp.text_frame
    tf_comp.word_wrap = True
    
    cp = tf_comp.paragraphs[0]
    cp.text = "BEFORE vs. AFTER SKILLPROOF"
    cp.font.size = Pt(12)
    cp.font.bold = True
    cp.font.color.rgb = TEXT_DARK

    t_before = tf_comp.add_paragraph()
    t_before.text = "❌ BEFORE: "
    t_before.font.bold = True
    t_before.font.color.rgb = RGBColor(220, 38, 38)
    t_before.font.size = Pt(10.5)
    r_b = t_before.add_run()
    r_b.text = "Resumes with unproven claims ('Proficient in Distributed Systems'). Recruiters guess competency or spend $500/candidate on screening. Non-target students get rejected automatically."
    r_b.font.bold = False
    r_b.font.color.rgb = TEXT_MUTED

    t_after = tf_comp.add_paragraph()
    t_after.text = "✅ AFTER: "
    t_after.font.bold = True
    t_after.font.color.rgb = ACCENT_GREEN
    t_after.font.size = Pt(10.5)
    t_after.space_before = Pt(4)
    r_a = t_after.add_run()
    r_a.text = "Transparent portfolios showing 82.8/100 score, verified 14-month commit graph, 32 source files, AI skill mapping, and a verified faculty sign-off. Recruiter finds qualified talent in seconds."
    r_a.font.bold = False
    r_a.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 4: 2. Solution Workflow & Personas
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "2. Solution Architecture", "User Workflow & Tri-Persona Ecosystem")

    personas = [
        ("Student / Builder", "Input & Evidence Ownership", [
            ("1. Profile Ingestion", "Connects GitHub handle or pastes repository URL."),
            ("2. Automated Parsing", "Engine analyzes commits, languages, code volume & docs."),
            ("3. Privacy Controls", "Student toggles project visibility (Public vs. Private)."),
            ("4. Portfolio Link", "Shares verifiable portfolio link directly with recruiters.")
        ], PRIMARY),
        ("Mentor / Faculty", "Human-in-the-Loop Verification", [
            ("1. Code Review", "Inspects project code quality and student execution."),
            ("2. Qualitative Review", "Leaves structured feedback (Strengths & Improvements)."),
            ("3. Star Rating", "Awards 1 to 5 star rating on technical depth."),
            ("4. Skill Verification", "Checks off AI-suggested skills into permanent verified status.")
        ], ACCENT_GREEN),
        ("Recruiter / Manager", "Precision Talent Discovery", [
            ("1. Skill & Score Filter", "Searches candidate pool by required tech stack & min score."),
            ("2. Evidence Inspection", "Views interactive commit timeline & verified skill badges."),
            ("3. Anti-Gaming Review", "Verifies zero burst commit dumps or unmaintained forks."),
            ("4. Direct Outreach", "Connects directly with verified high-signal candidates.")
        ], ACCENT_AMBER)
    ]

    for i, (p_name, p_role, steps, p_color) in enumerate(personas):
        left = Inches(0.8 + i * 3.95)
        add_card(s4, left, Inches(1.6), Inches(3.8), Inches(5.3), bg_color=CARD_BG, border_color=CARD_BORDER)
        
        # Header banner inside card
        b_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.6), Inches(3.8), Inches(0.8))
        b_box.fill.solid()
        b_box.fill.fore_color.rgb = p_color
        b_box.line.fill.background()
        btf = b_box.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.text = p_name
        bp.font.size = Pt(14)
        bp.font.bold = True
        bp.font.color.rgb = TEXT_LIGHT
        bp.alignment = PP_ALIGN.CENTER
        bp2 = btf.add_paragraph()
        bp2.text = p_role
        bp2.font.size = Pt(9.5)
        bp2.font.color.rgb = RGBColor(224, 231, 255)
        bp2.alignment = PP_ALIGN.CENTER

        # Steps
        stb = s4.shapes.add_textbox(left + Inches(0.2), Inches(2.55), Inches(3.4), Inches(4.2))
        stf = stb.text_frame
        stf.word_wrap = True
        for s_idx, (stitle, sdesc) in enumerate(steps):
            p = stf.add_paragraph() if s_idx > 0 else stf.paragraphs[0]
            p.text = stitle
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = TEXT_DARK
            p.space_before = Pt(8) if s_idx > 0 else Pt(0)
            
            p_sub = stf.add_paragraph()
            p_sub.text = sdesc
            p_sub.font.size = Pt(9.5)
            p_sub.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 5: 3. Prototype & System Architecture
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "3. Prototype & Architecture", "Working Prototype: Decoupled Full-Stack Architecture")

    # Left Column: Layered System Diagram
    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.3), bg_color=CARD_BG, border_color=CARD_BORDER)
    tb_l = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.9))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "SYSTEM ARCHITECTURE STACK"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    arch_layers = [
        ("Frontend SPA", "Semantic HTML5, CSS3 Glassmorphism tokens, Chart.js for commit graphs, Lucide icons, responsive navigation."),
        ("FastAPI Backend", "Asynchronous Python 3.11 server with REST endpoints, Pydantic v2 schemas, automated OpenAPI documentation."),
        ("Real-Time Sync Engine", "Bidirectional WebSocket connection (`/ws/feed`) broadcasting real-time imports, reviews, and recalculations."),
        ("Evidence Engine", "Pure deterministic scoring service (30% consistency, 25% complexity, 20% README, 15% recency, 10% community)."),
        ("Dual AI / NLP Engine", "Rule-based regex & NLP tokenizer for skills taxonomy with optional OpenAI GPT fallback integration."),
        ("Data Persistence", "SQLAlchemy ORM with SQLite database (models: Student, Project, Skill, MentorFeedback).")
    ]
    for idx, (layer, desc) in enumerate(arch_layers):
        lp = tf_l.add_paragraph()
        lp.text = f"[{idx+1}] {layer}: "
        lp.font.size = Pt(10.5)
        lp.font.bold = True
        lp.font.color.rgb = TEXT_DARK
        lp.space_before = Pt(6)
        r = lp.add_run()
        r.text = desc
        r.font.bold = False
        r.font.color.rgb = TEXT_MUTED

    # Right Column: Prototype Working Screens & Demo Highlights
    add_card(s5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), bg_color=CARD_BG_TINT, border_color=CARD_BORDER)
    tb_r = s5.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.9))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_r = tf_r.paragraphs[0]
    p_r.text = "PROTOTYPE IMPLEMENTATION & SCREENS"
    p_r.font.size = Pt(14)
    p_r.font.bold = True
    p_r.font.color.rgb = TEXT_DARK

    proto_points = [
        ("Live Single-Page App", "Loaded at http://127.0.0.1:8000/ with zero build-step overhead, instant response times, and dark theme."),
        ("Candidate Portfolios", "Pre-seeded with 5 realistic developer profiles (scores 35.2 to 90.8) displaying dynamic Chart.js commit graphs."),
        ("Interactive Scoring Audit", "Visual progress bars breaking down every sub-metric so candidates and recruiters see exact calculation logic."),
        ("Recruiter Search Matrix", "Real-time client/server multi-skill filtering, minimum score slider, and sorting by evidence level."),
        ("Mentor Feedback Modal", "Interactive rating, written feedback, and 1-click skill verification badges."),
        ("GitHub Importer Modal", "Allows on-demand import of any public GitHub repository with live progress feedback.")
    ]
    for idx, (item, desc) in enumerate(proto_points):
        rp = tf_r.add_paragraph()
        rp.text = f"✔ {item}: "
        rp.font.size = Pt(10.5)
        rp.font.bold = True
        rp.font.color.rgb = ACCENT_GREEN
        rp.space_before = Pt(6)
        r = rp.add_run()
        r.text = desc
        r.font.bold = False
        r.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 6: 3. Prototype Proof of Concept - The Scoring Model
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "3. Model & Proof of Concept", "The Deterministic Scoring Model & Anti-Gaming Formula")

    # Table for 5 Scoring Pillars
    pillars_table_shape = s6.shapes.add_table(6, 4, Inches(0.8), Inches(1.6), Inches(7.6), Inches(3.2))
    table = pillars_table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(1.0)
    table.columns[2].width = Inches(2.8)
    table.columns[3].width = Inches(1.6)

    headers = ["Evaluation Pillar", "Weight", "Mathematical Metric", "Formula Criteria"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT

    rows_data = [
        ("Commit Consistency", "30 pts", "Active months / total span ratio + commit count", "Diminishing returns curve"),
        ("Project Complexity", "25 pts", "Languages count + file count + repo size (KB)", "Max pts for 3+ langs, 50+ files"),
        ("README Quality", "20 pts", "Character length + sections (install/usage/code)", "Heuristic markdown scan"),
        ("Recent Activity", "15 pts", "Days since last commit timestamp", "Full pts if <30 days; scales down"),
        ("Community Signal", "10 pts", "GitHub stars count + original repo check", "Up to 6 pts for stars, 4 for original")
    ]
    for i, row in enumerate(rows_data, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if i % 2 == 0 else CARD_BG_TINT
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.5)
                p.font.color.rgb = TEXT_DARK

    # Right Card: Anti-Gaming Engine
    add_card(s6, Inches(8.6), Inches(1.6), Inches(3.9), Inches(5.3), bg_color=CARD_BG, border_color=RGBColor(239, 68, 68))
    tb_ag = s6.shapes.add_textbox(Inches(8.8), Inches(1.8), Inches(3.5), Inches(4.9))
    tf_ag = tb_ag.text_frame
    tf_ag.word_wrap = True

    p = tf_ag.paragraphs[0]
    p.text = "ANTI-GAMING PENALTIES"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(220, 38, 38)

    anti_items = [
        ("Single-Day Dump (-15 pts)", "Detects repos with >10 commits where 80%+ occurred within a 24h window (discourages artificial sprint gaming)."),
        ("Bare Fork Detection (-30 pts)", "Detects repositories forked from another author that contain zero student-authored original commits (score capped at 10)."),
        ("Minimal Shell (-10 pts)", "Detects trivial repositories with <2KB code or only 1-2 placeholder files."),
        ("No Private Token Abuse", "Rate limits GitHub queries and validates HTTP response status with graceful fallbacks.")
    ]
    for h, desc in anti_items:
        ap = tf_ag.add_paragraph()
        ap.text = f"⚠ {h}: "
        ap.font.size = Pt(10)
        ap.font.bold = True
        ap.font.color.rgb = TEXT_DARK
        ap.space_before = Pt(8)
        r = ap.add_run()
        r.text = desc
        r.font.bold = False
        r.font.color.rgb = TEXT_MUTED

    # Bottom Proof of Concept Card
    add_card(s6, Inches(0.8), Inches(5.1), Inches(7.6), Inches(1.8), bg_color=RGBColor(238, 242, 255), border_color=PRIMARY_LIGHT)
    tb_poc = s6.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(7.2), Inches(1.6))
    tf_poc = tb_poc.text_frame
    tf_poc.word_wrap = True
    
    pp = tf_poc.paragraphs[0]
    pp.text = "VALIDATED TEST SUITE & EMPIRICAL PROOF"
    pp.font.size = Pt(11)
    pp.font.bold = True
    pp.font.color.rgb = PRIMARY

    p_runs = [
        "• 13 Automated Pytest Unit & Integration Tests pass with 100% green status (`backend/tests/`).",
        "• Tested against edge cases: empty repositories, zero-commit accounts, massive 1000+ commit codebases.",
        "• High-performing students (e.g. Marcus: 90.8/100) vs. low-effort burst portfolios (Jordan: 35.2/100) are objectively separated."
    ]
    for r_line in p_runs:
        pl = tf_poc.add_paragraph()
        pl.text = r_line
        pl.font.size = Pt(9.5)
        pl.font.color.rgb = TEXT_DARK
        pl.space_before = Pt(3)

    # =========================================================================
    # SLIDE 7: 4. Technology & Safety
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "4. Technology & Safety", "Engineering Standards, Data Integrity & Responsible AI")

    tech_safety_quadrants = [
        ("Technology Stack", [
            ("Backend Framework", "FastAPI (Python 3.11) with async routing & automatic OpenAPI schema."),
            ("Data Persistence", "SQLAlchemy 2.0 ORM with relational SQLite database engine."),
            ("Frontend Engineering", "Vanilla ES6 JavaScript, HTML5, custom CSS tokens, Chart.js, Lucide."),
            ("Real-Time Engine", "Native WebSockets (`/ws/feed`) for multi-client event broadcasting.")
        ], PRIMARY),
        ("Data Handling & Integrity", [
            ("Read-Only GitHub Ingestion", "Only requests public repository metadata via official GitHub REST API."),
            ("No Source Code Stored", "Stores only metadata (commit dates, language bytes, file count, README text)."),
            ("Deterministic Algorithms", "Scoring is 100% reproducible with mathematical formulas and no hidden weights."),
            ("Sanitized Inputs", "Pydantic v2 data validation prevents SQL injection & malformed payloads.")
        ], PRIMARY_LIGHT),
        ("Privacy & Security", [
            ("Granular Privacy Toggle", "Students can switch any project between Public and Private with one click."),
            ("Zero Credential Logging", "Never stores GitHub passwords; uses optional client-scoped read tokens."),
            ("CORS & Origin Hardening", "Configured CORS middleware restricts unauthorized cross-origin requests."),
            ("Rate Limit Throttling", "Built-in GitHub quota tracking prevents API exhaustion and IP blacklisting.")
        ], ACCENT_GREEN),
        ("Human Control & Governance", [
            ("Human-in-the-Loop", "AI suggests skills, but only human mentors can grant 'Verified' credential status."),
            ("Faculty Accountability", "Every verified skill records the mentor name, email, and timestamp audit trail."),
            ("Override & Feedback", "Candidates can contest flags, review breakdown logs, and submit fresh commits."),
            ("AI Hallucination Safeguard", "Rule-based skill extraction runs deterministically; OpenAI LLM is optional.")
        ], ACCENT_AMBER)
    ]

    for i, (title, items, border_col) in enumerate(tech_safety_quadrants):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.6 + row * 2.75)
        
        add_card(s7, left, top, Inches(5.8), Inches(2.6), bg_color=CARD_BG, border_color=border_col)
        tb = s7.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.4), Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = border_col

        for subhead, text in items:
            lp = tf.add_paragraph()
            lp.text = f"• {subhead}: "
            lp.font.size = Pt(10)
            lp.font.bold = True
            lp.font.color.rgb = TEXT_DARK
            lp.space_before = Pt(3)
            r = lp.add_run()
            r.text = text
            r.font.bold = False
            r.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 8: 5. Business & Market Strategy
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "5. Business & Market Opportunity", "Market Opportunity, Business Model & Future Scope")

    b_cols = [
        ("Target Customers", [
            ("Universities & Bootcamps", "Computer science departments seeking higher graduate placement rates."),
            ("Tech Companies & Startups", "Engineering teams needing faster, higher-accuracy junior and intern hiring."),
            ("EdTech & Learning Platforms", "Coursera, Udacity, LeetCode wanting to add verified portfolio outcomes."),
            ("Individual Developers", "Self-taught builders needing objective proof to land interviews.")
        ]),
        ("Business Model", [
            ("Freemium for Students", "Free GitHub import, evidence score, and shareable portfolio link."),
            ("B2B Recruiter SaaS ($299/mo)", "Access to advanced candidate search, verified score filters, direct messaging."),
            ("University Enterprise Licensing", "Cohort dashboards, mentor verification workflows, accreditation reports."),
            ("API-as-a-Service", "Background skill verification API for ATS platforms (Greenhouse, Lever).")
        ]),
        ("Market Opportunity & Future Scope", [
            ("Total Addressable Market", "$30B+ Global Recruitment Software & Technical Assessment Market."),
            ("Multi-Platform Ingestion", "Future roadmap: GitLab, Bitbucket, Kaggle notebooks, LeetCode progress."),
            ("Automated Test Runner Sandbox", "Dockerized sandbox that actually compiles and runs student repo tests."),
            ("Decentralized Verifiable Credentials", "W3C-compliant digital credentials issued on-chain to prevent forgery.")
        ])
    ]

    for i, (title, items) in enumerate(b_cols):
        left = Inches(0.8 + i * 3.95)
        add_card(s8, left, Inches(1.6), Inches(3.8), Inches(5.3), bg_color=CARD_BG, border_color=CARD_BORDER)
        
        tb = s8.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(3.4), Inches(0.6))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        tb_body = s8.shapes.add_textbox(left + Inches(0.2), Inches(2.4), Inches(3.4), Inches(4.3))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        for idx, (subhead, text) in enumerate(items):
            bp = tf_body.add_paragraph() if idx > 0 else tf_body.paragraphs[0]
            bp.text = f"• {subhead}: "
            bp.font.size = Pt(10.5)
            bp.font.bold = True
            bp.font.color.rgb = TEXT_DARK
            bp.space_before = Pt(8) if idx > 0 else Pt(0)
            
            run = bp.add_run()
            run.text = text
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: 6. 7-Minute Presentation Timing & Summary
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "6. Presentation Roadmap", "7-Minute Delivery Script & Meeting Pacing Guide")

    # Timing table
    t_table_shape = s9.shapes.add_table(6, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.2))
    tt = t_table_shape.table
    tt.columns[0].width = Inches(1.8)
    tt.columns[1].width = Inches(1.2)
    tt.columns[2].width = Inches(3.2)
    tt.columns[3].width = Inches(5.5)

    th = ["Phase", "Time Window", "Speaker Focus / Action", "Key Message to Deliver"]
    for j, val in enumerate(th):
        c = tt.cell(0, j)
        c.text = val
        c.fill.solid()
        c.fill.fore_color.rgb = PRIMARY
        for p in c.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT

    t_rows = [
        ("1. Problem", "0:00 – 1:15 (75s)", "Hook audience with credential inflation crisis", "Keywords on resumes and course certificates are dead; hiring teams waste millions on screens."),
        ("2. Solution", "1:15 – 2:30 (75s)", "Introduce SkillProof & Before/After value", "SkillProof replaces self-reporting with git commit evidence, anti-gaming audits, and mentor verification."),
        ("3. Live Prototype", "2:30 – 4:30 (120s)", "Live UI walkthrough: Student -> Recruiter -> Mentor", "Show Marcus's 90.8 score, Chart.js timeline, recruiter skill search, and mentor 1-click verification."),
        ("4. Tech & Safety", "4:30 – 5:45 (75s)", "FastAPI stack, privacy toggle, human-in-loop", "Read-only metadata, privacy toggles for students, and mandatory human sign-off on AI-suggested skills."),
        ("5. Business Model", "5:45 – 7:00 (75s)", "Monetization, TAM, and strategic expansion", "$30B recruitment market, SaaS recruiter subscription, and university enterprise licensing.")
    ]
    for i, row in enumerate(t_rows, start=1):
        for j, val in enumerate(row):
            c = tt.cell(i, j)
            c.text = val
            c.fill.solid()
            c.fill.fore_color.rgb = CARD_BG if i % 2 == 0 else CARD_BG_TINT
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = TEXT_DARK
                if j == 1:
                    p.font.bold = True
                    p.font.color.rgb = PRIMARY

    # Callout bottom
    add_card(s9, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9), bg_color=RGBColor(238, 242, 255), border_color=PRIMARY_LIGHT)
    tb_b = s9.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.8))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    pb = tf_b.paragraphs[0]
    pb.text = "🎯 Q&A TRANSITION (Minute 7:00+): Open floor for examiner & stakeholder questions. (Refer to anticipated Q&A bank for defense on anti-gaming, rate limits, and verification)."
    pb.font.size = Pt(11)
    pb.font.bold = True
    pb.font.color.rgb = PRIMARY

    # Save presentation
    prs.save(output_path)
    print(f"[OK] Successfully generated presentation at: {output_path}")

if __name__ == "__main__":
    create_deck("SkillProof_Presentation.pptx")
